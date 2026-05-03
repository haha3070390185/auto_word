import os
import re
from docx import Document
from docx.shared import Pt, Cm, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from typing import List, Dict, Any, Optional, Tuple
from config import DEFAULT_FONT_NAME, DEFAULT_FONT_SIZE, DEFAULT_LINE_SPACING, DEFAULT_INDENTATION


class DocumentGenerator:
    def __init__(self, template_path: str):
        self.template_path = template_path
        self.placeholder_pattern = r'\{([^}]+)\}'
        self.placeholders = []
        self._extract_placeholders()

    def _extract_placeholders(self):
        try:
            doc = Document(self.template_path)
            self.placeholders = []
            
            for para in doc.paragraphs:
                matches = re.findall(self.placeholder_pattern, para.text)
                for match in matches:
                    if match not in self.placeholders:
                        self.placeholders.append(match)
            
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        for para in cell.paragraphs:
                            matches = re.findall(self.placeholder_pattern, para.text)
                            for match in matches:
                                if match not in self.placeholders:
                                    self.placeholders.append(match)
        
        except Exception as e:
            raise Exception(f"无法读取模板文件: {str(e)}")

    def get_placeholders(self) -> List[str]:
        return self.placeholders

    def _replace_in_paragraph(self, para, data: Dict[str, str]):
        for run in para.runs:
            for key, value in data.items():
                placeholder = '{' + key + '}'
                if placeholder in run.text:
                    run.text = run.text.replace(placeholder, value)

    def generate_document(self, data: Dict[str, str], output_path: str) -> str:
        try:
            doc = Document(self.template_path)
            
            for para in doc.paragraphs:
                self._replace_in_paragraph(para, data)
            
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        for para in cell.paragraphs:
                            self._replace_in_paragraph(para, data)
            
            for section in doc.sections:
                for para in section.header.paragraphs:
                    self._replace_in_paragraph(para, data)
                for para in section.footer.paragraphs:
                    self._replace_in_paragraph(para, data)
            
            doc.save(output_path)
            return output_path
        
        except Exception as e:
            raise Exception(f"生成文档失败: {str(e)}")

    def batch_generate(self, data_list: List[Dict[str, str]], 
                       output_dir: str,
                       filename_pattern: str = "文档_{index}.docx",
                       key_field: str = None) -> List[Tuple[str, Optional[str]]]:
        results = []
        
        if not os.path.exists(output_dir):
            os.makedirs(output_dir)
        
        for index, data in enumerate(data_list):
            try:
                if key_field and key_field in data and data[key_field]:
                    safe_filename = re.sub(r'[\\/:*?"<>|]', '_', data[key_field])
                    filename = f"{safe_filename}.docx"
                else:
                    filename = filename_pattern.replace("{index}", str(index + 1))
                
                output_path = os.path.join(output_dir, filename)
                
                counter = 1
                while os.path.exists(output_path):
                    name, ext = os.path.splitext(filename)
                    output_path = os.path.join(output_dir, f"{name}_{counter}{ext}")
                    counter += 1
                
                self.generate_document(data, output_path)
                results.append((output_path, None))
            
            except Exception as e:
                results.append((None, str(e)))
        
        return results


class DocumentFormatter:
    def __init__(self, file_path: str):
        self.file_path = file_path
        self.doc = None
        self._load_document()

    def _load_document(self):
        try:
            self.doc = Document(self.file_path)
        except Exception as e:
            raise Exception(f"无法打开文档: {str(e)}")

    def set_font(self, font_name: str = DEFAULT_FONT_NAME, 
                 font_size: float = DEFAULT_FONT_SIZE,
                 bold: bool = None, italic: bool = None):
        for para in self.doc.paragraphs:
            for run in para.runs:
                run.font.name = font_name
                run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
                run.font.size = Pt(font_size)
                if bold is not None:
                    run.font.bold = bold
                if italic is not None:
                    run.font.italic = italic
        
        for table in self.doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        for run in para.runs:
                            run.font.name = font_name
                            run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
                            run.font.size = Pt(font_size)
                            if bold is not None:
                                run.font.bold = bold
                            if italic is not None:
                                run.font.italic = italic

    def set_line_spacing(self, line_spacing: float = DEFAULT_LINE_SPACING):
        for para in self.doc.paragraphs:
            para.paragraph_format.line_spacing = line_spacing
        
        for table in self.doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        para.paragraph_format.line_spacing = line_spacing

    def set_indentation(self, indentation: float = DEFAULT_INDENTATION):
        indent_cm = Cm(indentation)
        for para in self.doc.paragraphs:
            para.paragraph_format.first_line_indent = indent_cm
        
        for table in self.doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        para.paragraph_format.first_line_indent = indent_cm

    def set_alignment(self, alignment: str = "left"):
        alignment_map = {
            "left": WD_ALIGN_PARAGRAPH.LEFT,
            "center": WD_ALIGN_PARAGRAPH.CENTER,
            "right": WD_ALIGN_PARAGRAPH.RIGHT,
            "justify": WD_ALIGN_PARAGRAPH.JUSTIFY
        }
        
        para_alignment = alignment_map.get(alignment.lower(), WD_ALIGN_PARAGRAPH.LEFT)
        
        for para in self.doc.paragraphs:
            para.paragraph_format.alignment = para_alignment
        
        for table in self.doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        para.paragraph_format.alignment = para_alignment

    def set_margins(self, top: float = 2.54, bottom: float = 2.54,
                    left: float = 3.17, right: float = 3.17):
        for section in self.doc.sections:
            section.top_margin = Cm(top)
            section.bottom_margin = Cm(bottom)
            section.left_margin = Cm(left)
            section.right_margin = Cm(right)

    def apply_formatting(self, font_name: str = None, font_size: float = None,
                         bold: bool = None, italic: bool = None,
                         line_spacing: float = None, indentation: float = None,
                         alignment: str = None, margins: Dict[str, float] = None):
        if font_name or font_size or bold is not None or italic is not None:
            self.set_font(
                font_name or DEFAULT_FONT_NAME,
                font_size or DEFAULT_FONT_SIZE,
                bold, italic
            )
        
        if line_spacing:
            self.set_line_spacing(line_spacing)
        
        if indentation:
            self.set_indentation(indentation)
        
        if alignment:
            self.set_alignment(alignment)
        
        if margins:
            self.set_margins(
                margins.get('top', 2.54),
                margins.get('bottom', 2.54),
                margins.get('left', 3.17),
                margins.get('right', 3.17)
            )

    def save(self, output_path: str = None):
        save_path = output_path or self.file_path
        self.doc.save(save_path)
        return save_path


class BatchDocumentFormatter:
    @staticmethod
    def format_files(file_paths: List[str], output_dir: str,
                     font_name: str = None, font_size: float = None,
                     bold: bool = None, italic: bool = None,
                     line_spacing: float = None, indentation: float = None,
                     alignment: str = None, margins: Dict[str, float] = None) -> List[Tuple[str, Optional[str]]]:
        results = []
        
        if not os.path.exists(output_dir):
            os.makedirs(output_dir)
        
        for file_path in file_paths:
            try:
                filename = os.path.basename(file_path)
                output_path = os.path.join(output_dir, filename)
                
                counter = 1
                while os.path.exists(output_path):
                    name, ext = os.path.splitext(filename)
                    output_path = os.path.join(output_dir, f"{name}_formatted{ext}")
                    counter += 1
                
                formatter = DocumentFormatter(file_path)
                formatter.apply_formatting(
                    font_name=font_name,
                    font_size=font_size,
                    bold=bold,
                    italic=italic,
                    line_spacing=line_spacing,
                    indentation=indentation,
                    alignment=alignment,
                    margins=margins
                )
                formatter.save(output_path)
                results.append((output_path, None))
            
            except Exception as e:
                results.append((file_path, str(e)))
        
        return results
