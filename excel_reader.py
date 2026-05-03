import openpyxl
from typing import List, Dict, Any, Optional
from openpyxl.utils import get_column_letter


class ExcelReader:
    def __init__(self, file_path: str):
        self.file_path = file_path
        self.workbook = None
        self.active_sheet = None
        self.headers = []
        self._load_workbook()

    def _load_workbook(self):
        try:
            self.workbook = openpyxl.load_workbook(self.file_path, data_only=True)
            self.active_sheet = self.workbook.active
            self._get_headers()
        except Exception as e:
            raise Exception(f"无法打开Excel文件: {str(e)}")

    def _get_headers(self):
        if self.active_sheet is None:
            return
        
        self.headers = []
        max_column = self.active_sheet.max_column
        
        for col in range(1, max_column + 1):
            cell_value = self.active_sheet.cell(row=1, column=col).value
            if cell_value is not None:
                self.headers.append(str(cell_value).strip())
            else:
                self.headers.append(f"列{col}")

    def get_headers(self) -> List[str]:
        return self.headers

    def get_sheet_names(self) -> List[str]:
        if self.workbook:
            return self.workbook.sheetnames
        return []

    def set_active_sheet(self, sheet_name: str):
        if self.workbook and sheet_name in self.workbook.sheetnames:
            self.active_sheet = self.workbook[sheet_name]
            self._get_headers()

    def read_all_data(self) -> List[Dict[str, Any]]:
        if self.active_sheet is None:
            return []
        
        data = []
        max_row = self.active_sheet.max_row
        max_col = self.active_sheet.max_column
        
        for row_num in range(2, max_row + 1):
            row_data = {}
            for col_num in range(1, max_col + 1):
                header = self.headers[col_num - 1] if col_num - 1 < len(self.headers) else f"列{col_num}"
                cell_value = self.active_sheet.cell(row=row_num, column=col_num).value
                
                if cell_value is not None:
                    if isinstance(cell_value, (int, float)):
                        if cell_value == int(cell_value):
                            row_data[header] = str(int(cell_value))
                        else:
                            row_data[header] = str(cell_value)
                    else:
                        row_data[header] = str(cell_value).strip()
                else:
                    row_data[header] = ""
            
            if any(row_data.values()):
                data.append(row_data)
        
        return data

    def read_data_by_range(self, start_row: int = 2, end_row: Optional[int] = None,
                            start_col: int = 1, end_col: Optional[int] = None) -> List[Dict[str, Any]]:
        if self.active_sheet is None:
            return []
        
        max_row = self.active_sheet.max_row
        max_col = self.active_sheet.max_column
        
        if end_row is None:
            end_row = max_row
        if end_col is None:
            end_col = max_col
        
        data = []
        for row_num in range(start_row, end_row + 1):
            row_data = {}
            for col_num in range(start_col, end_col + 1):
                header = self.headers[col_num - 1] if col_num - 1 < len(self.headers) else f"列{col_num}"
                cell_value = self.active_sheet.cell(row=row_num, column=col_num).value
                
                if cell_value is not None:
                    if isinstance(cell_value, (int, float)):
                        if cell_value == int(cell_value):
                            row_data[header] = str(int(cell_value))
                        else:
                            row_data[header] = str(cell_value)
                    else:
                        row_data[header] = str(cell_value).strip()
                else:
                    row_data[header] = ""
            
            if any(row_data.values()):
                data.append(row_data)
        
        return data

    def get_row_count(self) -> int:
        if self.active_sheet is None:
            return 0
        return self.active_sheet.max_row - 1

    def close(self):
        if self.workbook:
            self.workbook.close()
